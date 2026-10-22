from __future__ import annotations

import hashlib
import json
import re
from io import BytesIO
from pathlib import Path
from typing import Any

import streamlit as st
import yaml
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from workbench.execution import ExecutionService
from workbench.models import (
    InputDefinition,
    ParameterDefinition,
    ProcessorResult,
    ToolDefinition,
    WorkflowStepDefinition,
    content_types_compatible,
)
from workbench.operations import EDITOR_OPERATION_BY_ID, EDITOR_OPERATIONS
from workbench.processors import (
    compact_json_text,
    pretty_json_text,
    remove_json_string_character_text,
    transform_json_unicode_text,
)
from workbench.registry import ToolRegistry
from workbench.storage import WorkbenchStorage
from workbench.ui import apply_design_system, page_header, section_header, tool_header

ROOT = Path(__file__).resolve().parent
STORAGE = WorkbenchStorage(ROOT / "data")
REGISTRY = ToolRegistry([ROOT / "registry" / "tools.yaml", STORAGE.custom_tools_path])
EXECUTION = ExecutionService(REGISTRY, STORAGE)

JSON_EDITOR_TOOL_IDS = (
    "json_pretty",
    "json_compact",
    "json_show_unicode",
    "json_escape_unicode",
    "json_unwrap_one_layer",
    "text_remove_backslashes",
    "json_remove_newlines",
    "json_remove_carriage_returns",
    "json_remove_tabs",
    "json_remove_spaces",
)


def json_workbench_tool_ids() -> set[str]:
    tool_ids = set(JSON_EDITOR_TOOL_IDS)
    tool_ids.update(
        tool.id
        for tool in REGISTRY.all()
        if len(tool.inputs) == 1 and tool.inputs[0].content_type == "json"
    )
    return tool_ids


st.set_page_config(page_title="JsonProcessor", page_icon="🧰", layout="wide")
apply_design_system()


def render_parameter(param: ParameterDefinition, default: Any, key_prefix: str) -> Any:
    key = f"{key_prefix}:{param.id}"
    if param.type == "boolean":
        return st.checkbox(param.label, value=bool(default), help=param.help or None, key=key)
    if param.type == "integer":
        return st.number_input(
            param.label, value=int(default or 0), step=1, help=param.help or None, key=key
        )
    if param.type == "select":
        index = param.options.index(default) if default in param.options else 0
        return st.selectbox(
            param.label, param.options, index=index, help=param.help or None, key=key
        )
    if param.type == "textarea":
        return st.text_area(
            param.label, value=str(default or ""), height=180, help=param.help or None, key=key
        )
    return st.text_input(param.label, value=str(default or ""), help=param.help or None, key=key)


def transform_json_editor(
    editor_key: str,
    message_key: str,
    action: str,
    history_key: str | None = None,
) -> None:
    raw = str(st.session_state.get(editor_key, ""))
    if not raw.strip():
        st.session_state[message_key] = ("warning", "请先粘贴 JSON 内容")
        return
    try:
        compact = "\n" not in raw and "\r" not in raw
        if action == "pretty":
            result = pretty_json_text(raw)
        elif action == "compact":
            result = compact_json_text(raw)
        elif action == "remove_spaces":
            result = remove_json_string_character_text(raw, " ", compact=compact)
        elif action == "remove_newlines":
            result = remove_json_string_character_text(raw, "\n", compact=compact)
        elif action == "remove_carriage_returns":
            result = remove_json_string_character_text(raw, "\r", compact=compact)
        elif action == "remove_tabs":
            result = remove_json_string_character_text(raw, "\t", compact=compact)
        elif action == "unwrap_one_layer":
            decoded = json.loads(raw.strip())
            if not isinstance(decoded, str):
                raise TypeError("当前内容不是由双引号包裹的 JSON 字符串")
            result = decoded
        elif action == "remove_backslashes":
            result = raw.replace("\\", "")
        elif action == "show_unicode":
            result = transform_json_unicode_text(raw, ensure_ascii=False)
        elif action == "escape_unicode":
            result = transform_json_unicode_text(raw, ensure_ascii=True)
        else:
            raise ValueError(f"未知操作：{action}")
    except (TypeError, ValueError) as exc:
        st.session_state[message_key] = ("error", str(exc))
        return
    if history_key is not None and result != raw:
        history = list(st.session_state.get(history_key, []))
        history.append(raw)
        st.session_state[history_key] = history[-30:]
    st.session_state[editor_key] = result
    st.session_state[message_key] = ("success", "操作完成，可继续点击其他操作")


def render_json_toolbar(
    editor_key: str,
    actions: list[str],
    *,
    history_key: str | None = None,
) -> None:
    message_key = f"message:{editor_key}"
    definitions = {
        operation.id: (
            operation.label,
            operation.warning or operation.description,
        )
        for operation in EDITOR_OPERATIONS
    }
    st.markdown("##### 快捷操作")
    for offset in range(0, len(actions), 3):
        chunk = actions[offset : offset + 3]
        columns = st.columns(len(chunk))
        for column, action in zip(columns, chunk, strict=True):
            label, help_text = definitions[action]
            column.button(
                label,
                width="stretch",
                help=help_text,
                on_click=transform_json_editor,
                args=(editor_key, message_key, action, history_key),
                key=f"action:{editor_key}:{action}",
            )
    message = st.session_state.get(message_key)
    if message:
        level, text = message
        getattr(st, level)(text)


def render_input(
    tool: ToolDefinition, input_def: InputDefinition, key_prefix: str
) -> tuple[str, bytes] | None:
    mode = "上传文件"
    if input_def.allow_paste:
        mode = st.radio(
            f"{input_def.label} · 输入方式",
            ["直接粘贴", "上传文件"],
            horizontal=True,
            key=f"input-mode:{tool.id}:{input_def.id}",
        )

    if mode == "上传文件":
        uploaded = st.file_uploader(
            input_def.label,
            type=input_def.accept or None,
            help=input_def.help or None,
            key=f"upload:{tool.id}:{input_def.id}",
        )
        if uploaded is None:
            return None
        return uploaded.name, uploaded.getvalue()

    editor_key = f"pasted:{tool.id}:{input_def.id}"
    pasted = st.text_area(
        input_def.paste_label,
        height=280,
        placeholder=input_def.placeholder
        or ("在这里粘贴 JSON…" if input_def.content_type.startswith("json") else "在这里粘贴文本…"),
        key=editor_key,
    )
    if input_def.editor_actions:
        render_json_toolbar(editor_key, input_def.editor_actions)
        pasted = str(st.session_state.get(editor_key, pasted))
    if not pasted.strip():
        return None

    if input_def.content_type in {"json", "json_array"}:
        try:
            normalized = pretty_json_text(pasted)
            if input_def.content_type == "json_array" and not isinstance(json.loads(pasted), list):
                raise TypeError("输入 JSON 必须是数组")
        except (TypeError, ValueError) as exc:
            st.error(str(exc))
            return None
        else:
            with st.expander("即时格式化预览", expanded=bool(input_def.editor_actions)):
                st.code(normalized[:20_000], language="json")
                if len(normalized) > 20_000:
                    st.caption("内容较长，预览仅显示前 20,000 个字符。")
    suffix = "json" if input_def.content_type.startswith("json") else "txt"
    return f"pasted.{suffix}", pasted.encode("utf-8")


def render_inputs(tool: ToolDefinition, key_prefix: str) -> dict[str, tuple[str, bytes]] | None:
    sources: dict[str, tuple[str, bytes]] = {}
    for input_def in tool.inputs:
        if len(tool.inputs) > 1:
            st.markdown(f"#### 输入 · {input_def.label}")
        source = render_input(tool, input_def, key_prefix)
        if source is not None:
            sources[input_def.id] = source
    return sources if len(sources) == len(tool.inputs) else None


@st.cache_data(show_spinner=False)
def inspect_workbook(content: bytes) -> dict[str, Any]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    sheets: dict[str, list[list[Any]]] = {}
    max_columns: dict[str, int] = {}
    for worksheet in workbook.worksheets:
        rows = [list(row) for row in worksheet.iter_rows(max_row=12, values_only=True)]
        sheets[worksheet.title] = rows
        max_columns[worksheet.title] = max((len(row) for row in rows), default=1)
    result = {
        "active": workbook.active.title,
        "sheets": sheets,
        "max_columns": max_columns,
    }
    workbook.close()
    return result


def render_inspected_parameter(
    param: ParameterDefinition,
    default: Any,
    key_prefix: str,
    params: dict[str, Any],
    workbook_info: dict[str, Any],
) -> Any:
    key = f"{key_prefix}:{param.id}"
    if param.ui_control == "workbook_sheet":
        options = list(workbook_info["sheets"])
        preferred = str(default or workbook_info["active"])
        index = options.index(preferred) if preferred in options else 0
        return st.selectbox(param.label, options, index=index, key=f"{key}:detected")
    if param.ui_control == "workbook_column":
        sheet_name = str(params.get(param.sheet_parameter) or workbook_info["active"])
        rows = workbook_info["sheets"][sheet_name]
        header_row = int(params.get(param.header_row_parameter, 1))
        header = rows[header_row - 1] if 0 < header_row <= len(rows) else []
        count = max(workbook_info["max_columns"][sheet_name], len(header), 1)
        options = [get_column_letter(index) for index in range(1, count + 1)]
        if not param.required:
            options = ["", *options]
        labels = {"": "不写入 ID"}
        for index, letter in enumerate(options):
            if not letter:
                continue
            header_index = index - (1 if not param.required else 0)
            value = header[header_index] if header_index < len(header) else None
            labels[letter] = f"{letter} · {value}" if value not in (None, "") else letter
        preferred = str(default or "")
        index = options.index(preferred) if preferred in options else 0
        return st.selectbox(
            param.label,
            options,
            index=index,
            format_func=labels.get,
            key=f"{key}:detected",
        )
    return render_parameter(param, default, key_prefix)


def inspect_declared_inputs(
    tool: ToolDefinition,
    sources: dict[str, tuple[str, bytes]] | None,
) -> dict[str, dict[str, Any]]:
    if sources is None:
        return {}
    inspected: dict[str, dict[str, Any]] = {}
    for input_definition in tool.inputs:
        if input_definition.inspector != "xlsx":
            continue
        try:
            inspected[input_definition.id] = inspect_workbook(sources[input_definition.id][1])
        except Exception as exc:  # noqa: BLE001 - uploaded workbooks can fail in many ways
            st.error(f"{input_definition.label}无法读取：{exc}")
    return inspected


def render_result(result: ProcessorResult, run_id: str) -> None:
    st.success(f"{result.summary}（运行 ID：{run_id}）")
    if result.preview:
        with st.expander("结果预览", expanded=True):
            st.dataframe(result.preview, width="stretch", hide_index=True)

    st.markdown("#### 输出产物")
    shown_paths: set[str] = set()
    for artifact_id, artifact in result.artifacts.items():
        if artifact.path in shown_paths:
            continue
        shown_paths.add(artifact.path)
        output_path = Path(artifact.path)
        if not output_path.exists():
            st.warning(f"产物已不存在：{artifact.name}")
            continue
        left, right = st.columns([4, 1])
        left.write(f"**{artifact_id}** · {output_path.name} · {artifact.size:,} bytes")
        right.download_button(
            "下载",
            data=output_path.read_bytes(),
            file_name=output_path.name,
            mime=artifact.media_type,
            key=f"download:{run_id}:{artifact_id}",
            width="stretch",
        )
        if output_path.suffix == ".json":
            with st.expander(f"查看 {output_path.name}"):
                st.code(output_path.read_text(encoding="utf-8")[:20_000], language="json")


def decode_uploaded_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("无法识别文本编码，请转换为 UTF-8 后重试")


def undo_json_editor(editor_key: str, history_key: str, message_key: str) -> None:
    history = list(st.session_state.get(history_key, []))
    if not history:
        st.session_state[message_key] = ("warning", "没有可撤销的操作")
        return
    st.session_state[editor_key] = history.pop()
    st.session_state[history_key] = history
    st.session_state[message_key] = ("success", "已撤销上一步操作")


def clear_json_editor(editor_key: str, history_key: str, message_key: str) -> None:
    current = str(st.session_state.get(editor_key, ""))
    if current:
        history = list(st.session_state.get(history_key, []))
        history.append(current)
        st.session_state[history_key] = history[-30:]
    st.session_state[editor_key] = ""
    st.session_state[message_key] = ("success", "编辑器已清空，可撤销恢复")


def render_runner(
    tools: list[ToolDefinition],
    *,
    title: str,
    caption: str,
    horizontal_selector: bool = False,
) -> None:
    page_header("PROCESSING", title, caption)

    options = [tool.id for tool in tools]
    format_tool = lambda tool_id: next(tool.name for tool in tools if tool.id == tool_id)
    if horizontal_selector:
        selected_id = st.radio(
            "选择功能",
            options,
            format_func=format_tool,
            horizontal=True,
            key="json-text-tool-selector",
        )
    else:
        selected_id = st.selectbox(
            "选择脚本或工作流",
            options=options,
            format_func=format_tool,
            key="data-tool-selector",
        )
    tool = REGISTRY.get(selected_id)
    with st.container(border=True):
        tool_header(tool)

    presets = STORAGE.load_presets().get(tool.id, {}) if tool.parameters else {}
    selected_preset = "不使用预设"
    if tool.parameters:
        preset_names = [selected_preset, *sorted(presets)]
        selected_preset = st.selectbox("参数预设", preset_names)
    preset_values = presets.get(selected_preset, {})
    state_prefix = f"{tool.id}:{selected_preset}"

    input_column, parameter_column = st.columns([1.05, 0.95], gap="large")
    with input_column, st.container(border=True):
        section_header(1, "准备输入", "所有工具使用一致的上传与粘贴入口。")
        sources = render_inputs(tool, state_prefix)
        inspected_inputs = inspect_declared_inputs(tool, sources)

    params: dict[str, Any] = {}
    with parameter_column, st.container(border=True):
        section_header(2, "设置参数", "参数来自工具定义，可保存为可复用预设。")
        if not tool.parameters:
            st.info("这个原子操作不需要参数。")
        for param in tool.parameters:
            default = preset_values.get(param.id, param.default)
            workbook_info = inspected_inputs.get(param.source_input)
            if workbook_info is not None and param.ui_control != "default":
                params[param.id] = render_inspected_parameter(
                    param, default, state_prefix, params, workbook_info
                )
            else:
                params[param.id] = render_parameter(param, default, state_prefix)

        if tool.parameters:
            with st.expander("保存参数预设"):
                preset_name = st.text_input("预设名称", key=f"preset-name:{tool.id}")
                save_label = "更新预设" if preset_name.strip() in presets else "保存预设"
                if st.button(save_label, key=f"save-preset:{tool.id}"):
                    if not preset_name.strip():
                        st.error("请输入预设名称")
                    else:
                        STORAGE.save_preset(tool.id, preset_name.strip(), params)
                        st.success(f"已保存预设：{preset_name.strip()}")
            if selected_preset != "不使用预设" and st.button(
                "删除当前预设", key=f"delete-preset:{tool.id}:{selected_preset}"
            ):
                STORAGE.delete_preset(tool.id, selected_preset)
                st.rerun()

    for input_id, workbook_info in inspected_inputs.items():
        selected_sheet = next(
            (
                str(params[param.id])
                for param in tool.parameters
                if param.source_input == input_id
                and param.ui_control == "workbook_sheet"
                and params.get(param.id)
            ),
            workbook_info["active"],
        )
        with st.expander(f"{selected_sheet} · Excel 模板预览"):
            preview_rows = workbook_info["sheets"][selected_sheet]
            if preview_rows:
                st.dataframe(preview_rows, width="stretch", hide_index=True)
            else:
                st.info("该工作表当前为空。")

    with st.container(border=True):
        section_header(3, "运行并查看结果", "一次运行只执行当前脚本或工作流。")
        if sources is None:
            st.caption("完成全部输入并通过格式校验后即可运行。")
        if st.button("开始运行", type="primary", disabled=sources is None, width="stretch"):
            assert sources is not None
            try:
                with st.spinner("正在处理…"):
                    run_id, result = EXECUTION.run_tool(tool, sources, params)
                st.session_state["last_result"] = {
                    "run_id": run_id,
                    "tool_id": tool.id,
                    "result": result.model_dump(),
                }
            except Exception as exc:  # noqa: BLE001 - UI boundary must report processor errors
                st.error(f"运行失败：{type(exc).__name__}: {exc}")

        last_result = st.session_state.get("last_result")
        if last_result and last_result.get("tool_id") == tool.id:
            result = ProcessorResult.model_validate(last_result["result"])
            render_result(result, str(last_result["run_id"]))


def render_shared_json_operation(
    tool: ToolDefinition,
    current: str,
    json_is_valid: bool,
) -> None:
    presets = STORAGE.load_presets().get(tool.id, {}) if tool.parameters else {}
    selected_preset = "不使用预设"
    if tool.parameters:
        selected_preset = st.selectbox(
            "参数预设",
            [selected_preset, *sorted(presets)],
            key=f"shared-json-preset:{tool.id}",
        )
    preset_values = presets.get(selected_preset, {})
    state_prefix = f"shared-json:{tool.id}:{selected_preset}"
    params: dict[str, Any] = {}
    for param in tool.parameters:
        params[param.id] = render_parameter(
            param,
            preset_values.get(param.id, param.default),
            state_prefix,
        )

    if tool.parameters:
        with st.expander("保存当前参数为预设"):
            preset_name = st.text_input("预设名称", key=f"shared-preset-name:{tool.id}")
            if st.button("保存预设", key=f"shared-save-preset:{tool.id}"):
                if preset_name.strip():
                    STORAGE.save_preset(tool.id, preset_name.strip(), params)
                    st.success(f"已保存预设：{preset_name.strip()}")
                else:
                    st.error("请输入预设名称")

    if st.button(
        f"运行 · {tool.name}",
        type="primary",
        disabled=not json_is_valid,
        width="stretch",
        key=f"shared-json-run:{tool.id}",
    ):
        try:
            with st.spinner("正在处理当前 JSON…"):
                run_id, result = EXECUTION.run_tool(
                    tool,
                    {tool.primary_input.id: ("current.json", current.encode("utf-8"))},
                    params,
                )
            st.session_state[f"shared-json-result:{tool.id}"] = {
                "run_id": run_id,
                "input_sha256": hashlib.sha256(current.encode("utf-8")).hexdigest(),
                "result": result.model_dump(),
            }
        except Exception as exc:  # noqa: BLE001 - UI boundary
            st.error(f"运行失败：{type(exc).__name__}: {exc}")

    saved_result = st.session_state.get(f"shared-json-result:{tool.id}")
    if saved_result:
        current_digest = hashlib.sha256(current.encode("utf-8")).hexdigest()
        if saved_result["input_sha256"] != current_digest:
            st.warning("当前编辑器内容已改变，下方结果来自修改前的 JSON。")
        render_result(
            ProcessorResult.model_validate(saved_result["result"]),
            str(saved_result["run_id"]),
        )


def render_json_text_tools() -> None:
    page_header(
        "JSON WORKBENCH",
        "JSON 处理工作台",
        "一份输入贯穿编辑、原子操作、文件输出和工作流；每次只执行一个明确动作。",
    )

    editor_key = "json-text-workbench-editor"
    history_key = "json-text-workbench-history"
    message_key = f"message:{editor_key}"
    editor_column, action_column = st.columns([1.65, 0.85], gap="large")
    with editor_column, st.container(border=True):
        section_header(1, "准备内容", "上传文件或直接粘贴，共用同一个编辑区。")
        uploaded = st.file_uploader(
            "载入 JSON 或文本文件",
            type=["json", "txt"],
            key="json-text-workbench-upload",
        )
        if uploaded is not None:
            content = uploaded.getvalue()
            upload_token = hashlib.sha256(content).hexdigest()
            if st.session_state.get("json-text-workbench-upload-token") != upload_token:
                try:
                    st.session_state[editor_key] = decode_uploaded_text(content)
                    st.session_state[history_key] = []
                    st.session_state["json-text-workbench-upload-token"] = upload_token
                    st.session_state[message_key] = ("success", f"已载入：{uploaded.name}")
                except ValueError as exc:
                    st.error(str(exc))
        st.text_area(
            "JSON / 文本编辑器",
            height=500,
            placeholder='在此粘贴 JSON，例如 {"id":1} 或 "{\\"id\\":1}"',
            key=editor_key,
        )

    current = str(st.session_state.get(editor_key, ""))
    json_is_valid = False
    with action_column, st.container(border=True):
        section_header(2, "选择原子操作", "一个操作只产生一种明确变化。")
        selected_action = st.selectbox(
            "文本操作",
            [operation.id for operation in EDITOR_OPERATIONS],
            format_func=lambda operation_id: EDITOR_OPERATION_BY_ID[operation_id].display_name,
            key="json-text-operation-selector",
        )
        operation = EDITOR_OPERATION_BY_ID[selected_action]
        st.caption(operation.description)
        if operation.warning:
            st.warning(operation.warning)
        st.button(
            "应用当前操作",
            type="primary",
            width="stretch",
            on_click=transform_json_editor,
            args=(editor_key, message_key, selected_action, history_key),
        )
        message = st.session_state.get(message_key)
        if message:
            level, text = message
            getattr(st, level)(text)

        undo_column, clear_column = st.columns(2)
        undo_column.button(
            "撤销",
            width="stretch",
            on_click=undo_json_editor,
            args=(editor_key, history_key, message_key),
        )
        clear_column.button(
            "清空",
            width="stretch",
            on_click=clear_json_editor,
            args=(editor_key, history_key, message_key),
        )

        if current.strip():
            try:
                parsed = json.loads(current)
            except json.JSONDecodeError as exc:
                st.warning(f"当前不是合法 JSON：第 {exc.lineno} 行、第 {exc.colno} 列，{exc.msg}")
                suffix = "txt"
                mime = "text/plain"
            else:
                json_is_valid = True
                value_type = (
                    "对象"
                    if isinstance(parsed, dict)
                    else "数组"
                    if isinstance(parsed, list)
                    else "值"
                )
                st.success(f"合法 JSON · 顶层类型：{value_type}")
                suffix = "json"
                mime = "application/json"
            st.download_button(
                "下载当前内容",
                data=current.encode("utf-8"),
                file_name=f"json_text_result.{suffix}",
                mime=mime,
                width="stretch",
            )

    output_tools = [
        tool
        for tool in REGISTRY.all()
        if tool.id not in JSON_EDITOR_TOOL_IDS
        and len(tool.inputs) == 1
        and tool.inputs[0].content_type == "json"
    ]
    if not output_tools:
        st.info("当前没有使用 JSON 单输入的输出工具。")
        return
    with st.container(border=True):
        section_header(3, "生成文件或运行工作流", "下方工具直接使用当前编辑器内容。")
        selected_output_id = st.selectbox(
            "选择文件输出或工作流",
            [tool.id for tool in output_tools],
            format_func=lambda tool_id: next(
                tool.name for tool in output_tools if tool.id == tool_id
            ),
            key="json-output-tool-selector",
        )
        selected_output_tool = next(tool for tool in output_tools if tool.id == selected_output_id)
        tool_header(selected_output_tool)
        if not json_is_valid:
            st.info("请先在上方准备合法 JSON。")
        render_shared_json_operation(selected_output_tool, current, json_is_valid)


def render_data_tools() -> None:
    render_runner(
        [
            tool
            for tool in REGISTRY.all()
            if tool.id not in json_workbench_tool_ids()
        ],
        title="文件转换与工作流",
        caption="JSON、TXT、XLSX 本地处理与可复用工作流。",
    )


def render_catalog() -> None:
    page_header(
        "LIBRARY",
        "脚本与工作流",
        "统一查看原子脚本、组合工作流及其输入输出契约。",
    )
    query = st.text_input("搜索", placeholder="按名称、ID 或标签搜索")
    catalog_tools = REGISTRY.all()
    all_tags = sorted({tag for tool in catalog_tools for tag in tool.tags})
    selected_tags = st.multiselect("标签", all_tags)
    filtered_tools = []
    for tool in catalog_tools:
        haystack = " ".join([tool.id, tool.name, tool.description, *tool.tags]).lower()
        if query.strip() and query.strip().lower() not in haystack:
            continue
        if selected_tags and not set(selected_tags).issubset(tool.tags):
            continue
        filtered_tools.append(tool)
    st.caption(f"显示 {len(filtered_tools)} / {len(catalog_tools)} 个工具")
    for tool in filtered_tools:
        if tool.kind == "workflow":
            icon = "🔗"
            scope_label = "工作流"
        else:
            icon = "⚛️"
            scope_label = "原子脚本"
        with st.expander(f"{icon} {tool.name} · v{tool.version}"):
            st.caption(scope_label)
            st.write(tool.description)
            st.code(tool.id)
            st.write("输入：" + "、".join(item.label for item in tool.inputs))
            st.write(
                "输出："
                + "、".join(
                    f"{output.label} ({output.content_type})" for output in tool.outputs
                )
            )
            st.write("参数：" + ("、".join(param.label for param in tool.parameters) or "无"))
            if tool.steps:
                st.write("步骤：" + " → ".join(f"{step.id} ({step.uses})" for step in tool.steps))
            st.write("标签：" + "、".join(tool.tags))

    with st.expander("开发者配置 · registry/tools.yaml"):
        st.write("所有工具定义均来自 YAML，可以直接使用 Git 管理和审查。")
        st.code((ROOT / "registry" / "tools.yaml").read_text(encoding="utf-8"), language="yaml")


def render_workflow_designer() -> None:
    page_header(
        "WORKFLOWS",
        "工作流设计器",
        "用原子脚本组合固定场景；步骤顺序、输入映射和参数均持久化保存。",
    )
    scripts = [
        tool
        for tool in REGISTRY.all()
        if tool.kind == "script" and len(tool.inputs) == 1
    ]
    if not scripts:
        st.warning("当前没有可用于线性编排的单输入脚本。")
        return

    saved_message = st.session_state.pop("workflow-designer-message", None)
    if saved_message:
        st.success(saved_message)

    st.markdown("### 1. 基本信息")
    left, right = st.columns(2)
    workflow_id = left.text_input(
        "工作流 ID",
        placeholder="例如 ids_to_request_xlsx",
        help="仅允许小写字母、数字和下划线。",
    ).strip()
    workflow_name = right.text_input("工作流名称", placeholder="例如 批量生成请求文件").strip()
    description = st.text_area("用途说明", placeholder="说明输入、处理过程和输出结果。", height=90)
    step_count = int(st.number_input("步骤数量", min_value=2, max_value=6, value=2, step=1))

    st.markdown("### 2. 选择步骤并固定参数")
    steps: list[WorkflowStepDefinition] = []
    first_tool: ToolDefinition | None = None
    previous_tool: ToolDefinition | None = None
    for index in range(step_count):
        step_number = index + 1
        with st.container(border=True):
            compatible_scripts = (
                scripts
                if previous_tool is None
                else [
                    tool
                    for tool in scripts
                    if content_types_compatible(
                        previous_tool.primary_output.content_type,
                        tool.primary_input.content_type,
                    )
                ]
            )
            if not compatible_scripts:
                st.error(
                    f"没有脚本可接收上一步的 {previous_tool.primary_output.content_type} 输出。"
                )
                break
            selected_id = st.selectbox(
                f"步骤 {step_number} · 脚本",
                [tool.id for tool in compatible_scripts],
                format_func=lambda tool_id: REGISTRY.get(tool_id).name,
                key=f"designer:tool:{index}",
            )
            selected_tool = REGISTRY.get(selected_id)
            if first_tool is None:
                first_tool = selected_tool
            st.caption(
                f"{selected_tool.primary_input.content_type} → "
                f"{selected_tool.primary_output.content_type} · {selected_tool.description}"
            )
            fixed_params: dict[str, Any] = {}
            if selected_tool.parameters:
                st.markdown("**固定参数**")
                for param in selected_tool.parameters:
                    fixed_params[param.id] = render_parameter(
                        param,
                        param.default,
                        f"designer:{index}:{selected_tool.id}",
                    )
            source_expression = (
                f"${{inputs.{first_tool.inputs[0].id}}}"
                if index == 0
                else f"${{steps.step_{index}.outputs.result}}"
            )
            steps.append(
                WorkflowStepDefinition(
                    id=f"step_{step_number}",
                    uses=selected_tool.id,
                    inputs={selected_tool.inputs[0].id: source_expression},
                    params=fixed_params,
                )
            )
            previous_tool = selected_tool

    candidate: ToolDefinition | None = None
    validation_error = ""
    if not re.fullmatch(r"[a-z][a-z0-9_]*", workflow_id):
        validation_error = "工作流 ID 必须以小写字母开头，只能包含小写字母、数字和下划线。"
    elif not workflow_name:
        validation_error = "请输入工作流名称。"
    elif not description.strip():
        validation_error = "请输入用途说明。"
    elif workflow_id in {tool.id for tool in REGISTRY.all()}:
        validation_error = "该工作流 ID 已存在，请使用新的 ID。"
    elif len(steps) != step_count:
        validation_error = "当前步骤之间没有兼容的数据类型，无法保存工作流。"
    elif first_tool is not None and previous_tool is not None and len(steps) == step_count:
        try:
            candidate = ToolDefinition(
                id=workflow_id,
                name=workflow_name,
                kind="workflow",
                description=description.strip(),
                version="1",
                tags=["自定义", "工作流"],
                inputs=[first_tool.inputs[0].model_copy(deep=True)],
                outputs=[previous_tool.primary_output.model_copy(deep=True)],
                parameters=[],
                steps=steps,
            )
        except ValueError as exc:
            validation_error = str(exc)

    st.markdown("### 3. 检查并保存")
    st.info("当前设计器按顺序把上一步的主要产物传给下一步，适合参数固定的重复场景。")
    if validation_error:
        st.warning(validation_error)
    if candidate is not None:
        yaml_text = yaml.safe_dump(
            {"tools": [candidate.model_dump(exclude_none=True)]},
            allow_unicode=True,
            sort_keys=False,
        )
        with st.expander("查看生成的 YAML"):
            st.code(yaml_text, language="yaml")
        left, right = st.columns(2)
        if left.button("保存到本机工具库", type="primary", width="stretch"):
            STORAGE.save_custom_tool(candidate)
            st.session_state["workflow-designer-message"] = f"已保存：{candidate.name}"
            st.rerun()
        right.download_button(
            "下载 YAML",
            data=yaml_text.encode("utf-8"),
            file_name=f"{candidate.id}.yaml",
            mime="text/yaml",
            width="stretch",
        )

    custom_tools = STORAGE.load_custom_tools()
    if custom_tools:
        st.divider()
        st.markdown("### 已保存的自定义工作流")
        selected_custom = st.selectbox(
            "选择工作流",
            [tool.id for tool in custom_tools],
            format_func=lambda tool_id: next(
                tool.name for tool in custom_tools if tool.id == tool_id
            ),
        )
        confirm_delete = st.checkbox("我确认删除这个自定义工作流")
        if st.button("删除工作流", disabled=not confirm_delete):
            STORAGE.delete_custom_tool(selected_custom)
            st.rerun()


def render_history() -> None:
    page_header("RUNS", "运行历史", "查看每次执行的输入、参数、步骤状态和全部输出产物。")
    runs = STORAGE.list_runs()
    if not runs:
        st.info("还没有运行记录。")
        return
    table = []
    for run in runs:
        table.append(
            {
                "运行 ID": run["id"],
                "工具": run["tool_name"],
                "类型": "工作流" if run["kind"] == "workflow" else "脚本",
                "状态": run["status"],
                "开始时间（UTC）": run["started_at"],
                "输入文件": run["input_name"],
                "摘要": run["summary"] or run["error"] or "",
            }
        )
    st.dataframe(table, width="stretch", hide_index=True)

    selected = st.selectbox("查看运行详情", [run["id"] for run in runs])
    run = next(item for item in runs if item["id"] == selected)
    details = dict(run)
    details["params"] = json.loads(details.pop("params_json"))
    st.json(details)
    inputs = STORAGE.list_run_inputs(selected)
    if inputs:
        st.markdown("#### 输入文件")
        st.dataframe(inputs, width="stretch", hide_index=True)
    steps = STORAGE.list_step_runs(selected)
    if steps:
        st.markdown("#### 步骤记录")
        st.dataframe(steps, width="stretch", hide_index=True)
    manifest = STORAGE.load_run_manifest(selected)
    if manifest and manifest.get("result"):
        artifacts = ProcessorResult.model_validate(manifest["result"])
        st.markdown("#### 全部运行产物")
        render_result(artifacts, selected)


def render_about() -> None:
    page_header("ABOUT", "关于", "本地优先、配置驱动、面向重复文件处理场景的开源工作台。")
    st.markdown(
        """
        JsonProcessor 是本地优先的开源文件处理工作台。当前版本不会把文件上传到外部服务。

        - 项目许可证：AGPL-3.0-or-later
        - 配置格式：YAML / JSON
        - 元数据：SQLite
        - 产物：普通 TXT / XLSX 文件
        """
    )


pages = {
    "处理": [
        st.Page(render_json_text_tools, title="JSON 工作台", icon="🧩", default=True),
        st.Page(render_data_tools, title="文件处理", icon="📁"),
    ],
    "构建": [
        st.Page(render_workflow_designer, title="工作流设计器", icon="🔗"),
        st.Page(render_catalog, title="脚本与工作流", icon="🧱"),
    ],
    "管理": [
        st.Page(render_history, title="运行历史", icon="🕘"),
        st.Page(render_about, title="关于", icon="ℹ️"),
    ],
}
navigation = st.navigation(pages)
with st.sidebar:
    st.markdown("### JsonProcessor")
    st.caption("Local processing workbench · v0.5.0")
navigation.run()
