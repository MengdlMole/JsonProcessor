from __future__ import annotations

import json
import re
from collections.abc import Callable
from pathlib import Path
from typing import Any

from jinja2 import StrictUndefined
from jinja2.sandbox import SandboxedEnvironment
from openpyxl import Workbook, load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from .models import ProcessorResult

_PATH_TOKEN = re.compile(r"([^.[\]]+)|\[(\*|\d+)\]")
_TEMPLATE_ENVIRONMENT = SandboxedEnvironment(undefined=StrictUndefined, autoescape=False)


def _read_json(path: Path) -> Any:
    raw = _read_text(path)
    try:
        return json.loads(raw)
    except json.JSONDecodeError as exc:
        raise ValueError(
            f"JSON 解析失败：第 {exc.lineno} 行、第 {exc.colno} 列，{exc.msg}"
        ) from exc


def compact_json_text(raw: str) -> str:
    """Remove only JSON layout whitespace, preserving every string token verbatim."""
    try:
        json.loads(raw.strip())
    except json.JSONDecodeError as exc:
        raise ValueError(
            f"JSON 解析失败：第 {exc.lineno} 行、第 {exc.colno} 列，{exc.msg}"
        ) from exc
    output: list[str] = []
    in_string = False
    escaped = False
    for character in raw.strip():
        if in_string:
            output.append(character)
            if escaped:
                escaped = False
            elif character == "\\":
                escaped = True
            elif character == '"':
                in_string = False
        elif character == '"':
            in_string = True
            output.append(character)
        elif not character.isspace():
            output.append(character)
    return "".join(output)


def pretty_json_text(raw: str, *, indent: int = 2) -> str:
    """Format JSON structure while preserving the spelling of string tokens."""
    compact = compact_json_text(raw)
    output: list[str] = []
    depth = 0
    in_string = False
    escaped = False
    matching = {"{": "}", "[": "]"}
    for index, character in enumerate(compact):
        if in_string:
            output.append(character)
            if escaped:
                escaped = False
            elif character == "\\":
                escaped = True
            elif character == '"':
                in_string = False
            continue
        if character == '"':
            in_string = True
            output.append(character)
        elif character in matching:
            output.append(character)
            if index + 1 < len(compact) and compact[index + 1] != matching[character]:
                depth += 1
                output.append("\n" + " " * (depth * indent))
        elif character in "}]":
            if index > 0 and compact[index - 1] not in "{[":
                depth -= 1
                output.append("\n" + " " * (depth * indent))
            output.append(character)
        elif character == ",":
            output.append(",\n" + " " * (depth * indent))
        elif character == ":":
            output.append(": ")
        else:
            output.append(character)
    return "".join(output)


def transform_json_unicode_text(raw: str, *, ensure_ascii: bool) -> str:
    """Change only JSON string token encoding and preserve layout outside strings."""
    try:
        json.loads(raw.strip())
    except json.JSONDecodeError as exc:
        raise ValueError(
            f"JSON 解析失败：第 {exc.lineno} 行、第 {exc.colno} 列，{exc.msg}"
        ) from exc
    output: list[str] = []
    index = 0
    while index < len(raw):
        if raw[index] != '"':
            output.append(raw[index])
            index += 1
            continue
        end = index + 1
        escaped = False
        while end < len(raw):
            character = raw[end]
            if escaped:
                escaped = False
            elif character == "\\":
                escaped = True
            elif character == '"':
                break
            end += 1
        token = raw[index : end + 1]
        output.append(json.dumps(json.loads(token), ensure_ascii=ensure_ascii))
        index = end + 1
    return "".join(output)


def _remove_character_from_strings(value: Any, character: str) -> Any:
    if isinstance(value, str):
        return value.replace(character, "")
    if isinstance(value, list):
        return [_remove_character_from_strings(item, character) for item in value]
    if isinstance(value, dict):
        return {key: _remove_character_from_strings(item, character) for key, item in value.items()}
    return value


def _count_character(value: Any, character: str) -> int:
    if isinstance(value, str):
        return value.count(character)
    if isinstance(value, list):
        return sum(_count_character(item, character) for item in value)
    if isinstance(value, dict):
        return sum(_count_character(item, character) for item in value.values())
    return 0


def remove_json_string_character_text(
    raw: str,
    character: str,
    *,
    compact: bool = False,
) -> str:
    try:
        value = json.loads(raw.strip())
    except json.JSONDecodeError as exc:
        raise ValueError(
            f"JSON 解析失败：第 {exc.lineno} 行、第 {exc.colno} 列，{exc.msg}"
        ) from exc
    cleaned = _remove_character_from_strings(value, character)
    if compact:
        return json.dumps(cleaned, ensure_ascii=False, separators=(",", ":"))
    return json.dumps(cleaned, ensure_ascii=False, indent=2)


def _read_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("无法识别文本编码，请转换为 UTF-8 后重试")


def _input_path(inputs: dict[str, Path], input_id: str) -> Path:
    try:
        return inputs[input_id]
    except KeyError as exc:
        if len(inputs) == 1:
            return next(iter(inputs.values()))
        raise ValueError(f"处理器缺少输入：{input_id}") from exc


def parse_path(expression: str) -> list[str | int]:
    expression = expression.strip().lstrip("$.")
    if not expression:
        return []
    tokens: list[str | int] = []
    cursor = 0
    for match in _PATH_TOKEN.finditer(expression):
        separator = expression[cursor : match.start()]
        if separator not in {"", "."}:
            raise ValueError(f"无效字段路径：{expression!r}")
        key, index = match.groups()
        token = key or index
        tokens.append(int(token) if token.isdigit() else token)
        cursor = match.end()
    if cursor != len(expression) or not tokens:
        raise ValueError(f"无效字段路径：{expression!r}")
    return tokens


def select_values(data: Any, expression: str) -> list[Any]:
    values = [data]
    for token in parse_path(expression):
        next_values: list[Any] = []
        for value in values:
            if token == "*":
                if isinstance(value, list):
                    next_values.extend(value)
                elif isinstance(value, dict):
                    next_values.extend(value.values())
            elif isinstance(token, int):
                if isinstance(value, list) and -len(value) <= token < len(value):
                    next_values.append(value[token])
            elif isinstance(value, dict) and token in value:
                next_values.append(value[token])
        values = next_values
    return values


def _deduplicate(values: list[Any]) -> list[Any]:
    result: list[Any] = []
    seen: set[str] = set()
    for value in values:
        marker = json.dumps(value, ensure_ascii=False, sort_keys=True)
        if marker not in seen:
            seen.add(marker)
            result.append(value)
    return result


def _safe_sheet_name(value: str) -> str:
    value = re.sub(r"[\\/*?:\[\]]", "_", value).strip()
    return (value or "Sheet1")[:31]


def _save_workbook(
    rows: list[list[Any]], headers: list[str], output_path: Path, sheet: str
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = _safe_sheet_name(sheet)
    worksheet.freeze_panes = "A2"
    worksheet.append([_excel_safe(value) for value in headers])
    for row in rows:
        worksheet.append([_excel_safe(value) for value in row])
    for index, header in enumerate(headers, start=1):
        sample_lengths = [len(str(header))]
        sample_lengths.extend(len(str(row[index - 1] or "")) for row in rows[:100])
        worksheet.column_dimensions[get_column_letter(index)].width = min(
            max(sample_lengths) + 2, 60
        )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def _excel_safe(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _require_json_array(path: Path) -> list[Any]:
    value = _read_json(path)
    if not isinstance(value, list):
        raise TypeError("输入 JSON 必须是数组")
    return value


def _write_json_value(
    value: Any,
    output_dir: Path,
    output_name: str,
    summary: str,
    preview: list[dict[str, Any]],
) -> ProcessorResult:
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / output_name
    output_path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    return ProcessorResult.from_path(
        output_path,
        summary=summary,
        preview=preview,
        media_type="application/json",
    )


def json_select_values(
    inputs: dict[str, Path], output_dir: Path, params: dict[str, Any]
) -> ProcessorResult:
    values = select_values(
        _read_json(_input_path(inputs, "source_json")), str(params["field_path"])
    )
    return _write_json_value(
        values,
        output_dir,
        "selected_values.json",
        f"已选择 {len(values)} 个值",
        [{"序号": index + 1, "值": value} for index, value in enumerate(values[:50])],
    )


def text_lines_to_json_array(
    inputs: dict[str, Path], output_dir: Path, params: dict[str, Any]
) -> ProcessorResult:
    values = _read_text(_input_path(inputs, "source_text")).splitlines()
    return _write_json_value(
        values,
        output_dir,
        "lines.json",
        f"已拆分 {len(values)} 行",
        [{"序号": index + 1, "值": value} for index, value in enumerate(values[:50])],
    )


def json_trim_string_values(
    inputs: dict[str, Path], output_dir: Path, params: dict[str, Any]
) -> ProcessorResult:
    values = _require_json_array(_input_path(inputs, "source_json"))
    trimmed = [value.strip() if isinstance(value, str) else value for value in values]
    return _write_json_value(
        trimmed,
        output_dir,
        "trimmed_values.json",
        f"已处理 {len(trimmed)} 个数组元素",
        [{"序号": index + 1, "值": value} for index, value in enumerate(trimmed[:50])],
    )


def json_filter_empty_values(
    inputs: dict[str, Path], output_dir: Path, params: dict[str, Any]
) -> ProcessorResult:
    values = _require_json_array(_input_path(inputs, "source_json"))
    filtered = [value for value in values if value not in (None, "")]
    return _write_json_value(
        filtered,
        output_dir,
        "non_empty_values.json",
        f"已移除 {len(values) - len(filtered)} 个空值",
        [{"序号": index + 1, "值": value} for index, value in enumerate(filtered[:50])],
    )


def json_deduplicate_values(
    inputs: dict[str, Path], output_dir: Path, params: dict[str, Any]
) -> ProcessorResult:
    values = _require_json_array(_input_path(inputs, "source_json"))
    unique = _deduplicate(values)
    return _write_json_value(
        unique,
        output_dir,
        "unique_values.json",
        f"已移除 {len(values) - len(unique)} 个重复值",
        [{"序号": index + 1, "值": value} for index, value in enumerate(unique[:50])],
    )


def json_values_to_lines(
    inputs: dict[str, Path], output_dir: Path, params: dict[str, Any]
) -> ProcessorResult:
    values = _require_json_array(_input_path(inputs, "source_json"))
    content = "\n".join(str(value) for value in values)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / "values.txt"
    output_path.write_text(content, encoding="utf-8")
    return ProcessorResult.from_path(
        output_path,
        summary=f"已写入 {len(values)} 行",
        preview=[{"文本结果": content[:10_000]}],
        media_type="text/plain",
    )


def json_values_to_quoted_list(
    inputs: dict[str, Path], output_dir: Path, params: dict[str, Any]
) -> ProcessorResult:
    values = _require_json_array(_input_path(inputs, "source_json"))
    content = ",".join(json.dumps(str(value), ensure_ascii=False) for value in values)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / "quoted_values.txt"
    output_path.write_text(content, encoding="utf-8")
    return ProcessorResult.from_path(
        output_path,
        summary=f"已序列化 {len(values)} 个值",
        preview=[{"文本结果": content[:10_000]}],
        media_type="text/plain",
    )


def json_render_template_rows(
    inputs: dict[str, Path], output_dir: Path, params: dict[str, Any]
) -> ProcessorResult:
    values = _require_json_array(_input_path(inputs, "source_json"))
    rendered_rows = _render_json_rows(values, params)
    column_name = str(params.get("column_name", "json_body"))
    rows = [{"id": row[0], column_name: row[1]} for row in rendered_rows]
    return _write_json_value(
        rows,
        output_dir,
        "rendered_rows.json",
        f"已渲染 {len(rows)} 条 JSON",
        rows[:50],
    )


def json_map_records(
    inputs: dict[str, Path], output_dir: Path, params: dict[str, Any]
) -> ProcessorResult:
    records = _require_json_array(_input_path(inputs, "source_json"))
    mapping = json.loads(str(params["column_mapping"]))
    if not isinstance(mapping, dict) or not mapping:
        raise ValueError("列映射必须是非空 JSON 对象")
    rows: list[dict[str, Any]] = []
    for record in records:
        row: dict[str, Any] = {}
        for header, path_expression in mapping.items():
            values = select_values(record, str(path_expression)) if path_expression else [record]
            row[str(header)] = values[0] if len(values) == 1 else values
        rows.append(row)
    return _write_json_value(
        rows,
        output_dir,
        "mapped_rows.json",
        f"已映射 {len(rows)} 行、{len(mapping)} 列",
        rows[:50],
    )


def json_rows_to_xlsx(
    inputs: dict[str, Path], output_dir: Path, params: dict[str, Any]
) -> ProcessorResult:
    records = _require_json_array(_input_path(inputs, "source_json"))
    if any(not isinstance(record, dict) for record in records):
        raise TypeError("写入 XLSX 的 JSON 数组元素必须全部是对象")
    headers = list(dict.fromkeys(str(key) for record in records for key in record))
    table_rows: list[list[Any]] = []
    preview: list[dict[str, Any]] = []
    for record in records:
        row_values: list[Any] = []
        preview_row: dict[str, Any] = {}
        for header in headers:
            value = record.get(header)
            if isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
            row_values.append(value)
            preview_row[header] = value
        table_rows.append(row_values)
        if len(preview) < 50:
            preview.append(preview_row)
    output_path = output_dir / "rows.xlsx"
    _save_workbook(table_rows, headers, output_path, str(params.get("sheet_name", "Data")))
    return ProcessorResult.from_path(
        output_path,
        summary=f"已写入 {len(table_rows)} 行、{len(headers)} 列",
        preview=preview,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _write_text_result(
    content: str,
    output_dir: Path,
    output_name: str,
    summary: str,
    preview_label: str,
) -> ProcessorResult:
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / output_name
    output_path.write_text(content, encoding="utf-8")
    return ProcessorResult.from_path(
        output_path,
        summary=summary,
        preview=[{preview_label: content[:10_000]}],
        media_type="application/json",
    )


def json_pretty(
    inputs: dict[str, Path], output_dir: Path, params: dict[str, Any]
) -> ProcessorResult:
    content = pretty_json_text(_read_text(_input_path(inputs, "source_json")))
    return _write_text_result(
        content, output_dir, "pretty.json", "已格式化 JSON 缩进", "格式化结果"
    )


def json_compact(
    inputs: dict[str, Path], output_dir: Path, params: dict[str, Any]
) -> ProcessorResult:
    content = compact_json_text(_read_text(_input_path(inputs, "source_json")))
    return _write_text_result(content, output_dir, "compact.json", "已压缩 JSON 排版", "压缩结果")


def json_show_unicode(
    inputs: dict[str, Path], output_dir: Path, params: dict[str, Any]
) -> ProcessorResult:
    content = transform_json_unicode_text(
        _read_text(_input_path(inputs, "source_json")), ensure_ascii=False
    )
    return _write_text_result(
        content, output_dir, "unicode_visible.json", "已显示 Unicode 字符", "转换结果"
    )


def json_escape_unicode(
    inputs: dict[str, Path], output_dir: Path, params: dict[str, Any]
) -> ProcessorResult:
    content = transform_json_unicode_text(
        _read_text(_input_path(inputs, "source_json")), ensure_ascii=True
    )
    return _write_text_result(
        content, output_dir, "unicode_escaped.json", "已转义 Unicode 字符", "转换结果"
    )


def json_unwrap_one_layer(
    inputs: dict[str, Path], output_dir: Path, params: dict[str, Any]
) -> ProcessorResult:
    raw = _read_text(_input_path(inputs, "source_text")).strip()
    try:
        decoded = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise ValueError(
            f"无法解码 JSON 字符串：第 {exc.lineno} 行、第 {exc.colno} 列，{exc.msg}"
        ) from exc
    if not isinstance(decoded, str):
        raise TypeError("输入必须是由双引号包裹的 JSON 字符串；本工具只解码一层")
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / "unwrapped.txt"
    output_path.write_text(decoded, encoding="utf-8")
    return ProcessorResult.from_path(
        output_path,
        summary="已解码一层 JSON 字符串",
        preview=[{"解码结果": decoded[:10_000]}],
        media_type="text/plain",
    )


def text_remove_backslashes(
    inputs: dict[str, Path], output_dir: Path, params: dict[str, Any]
) -> ProcessorResult:
    raw = _read_text(_input_path(inputs, "source_text"))
    removed_count = raw.count("\\")
    content = raw.replace("\\", "")
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / "without_backslashes.txt"
    output_path.write_text(content, encoding="utf-8")
    return ProcessorResult.from_path(
        output_path,
        summary=f"已删除 {removed_count} 个反斜杠",
        preview=[{"处理结果": content[:10_000]}],
        media_type="text/plain",
    )


def _json_remove_string_character(
    inputs: dict[str, Path],
    output_dir: Path,
    character: str,
    character_name: str,
    output_name: str,
) -> ProcessorResult:
    parsed = _read_json(_input_path(inputs, "source_json"))
    removed_count = _count_character(parsed, character)
    cleaned = _remove_character_from_strings(parsed, character)
    content = json.dumps(cleaned, ensure_ascii=False, indent=2)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / output_name
    output_path.write_text(content, encoding="utf-8")
    return ProcessorResult.from_path(
        output_path,
        summary=f"已移除 {removed_count} 个{character_name}",
        preview=[{"清理结果": content[:10_000]}],
        media_type="application/json",
    )


def json_remove_newlines(
    inputs: dict[str, Path], output_dir: Path, params: dict[str, Any]
) -> ProcessorResult:
    return _json_remove_string_character(inputs, output_dir, "\n", "换行符", "no_newlines.json")


def json_remove_carriage_returns(
    inputs: dict[str, Path], output_dir: Path, params: dict[str, Any]
) -> ProcessorResult:
    return _json_remove_string_character(
        inputs, output_dir, "\r", "回车符", "no_carriage_returns.json"
    )


def json_remove_tabs(
    inputs: dict[str, Path], output_dir: Path, params: dict[str, Any]
) -> ProcessorResult:
    return _json_remove_string_character(inputs, output_dir, "\t", "制表符", "no_tabs.json")


def json_remove_spaces(
    inputs: dict[str, Path], output_dir: Path, params: dict[str, Any]
) -> ProcessorResult:
    return _json_remove_string_character(inputs, output_dir, " ", "ASCII 空格", "no_spaces.json")


def _parse_common_params(raw: str) -> dict[str, Any]:
    value = json.loads(raw or "{}")
    if not isinstance(value, dict):
        raise TypeError("公共参数必须是 JSON 对象")
    return value


def _render_json_rows(ids: list[Any], params: dict[str, Any]) -> list[list[Any]]:
    template = _TEMPLATE_ENVIRONMENT.from_string(str(params["json_template"]))
    common = _parse_common_params(str(params.get("common_params", "{}")))
    id_variable = str(params.get("id_variable", "id"))
    rendered_rows: list[list[Any]] = []
    for value in ids:
        context = dict(common)
        context[id_variable] = value
        rendered = template.render(**context)
        parsed = json.loads(rendered)
        rendered_rows.append([value, json.dumps(parsed, ensure_ascii=False, separators=(",", ":"))])
    return rendered_rows


def _resolve_column(worksheet: Any, selector: str, header_row: int) -> int:
    value = selector.strip()
    if not value:
        raise ValueError("列名或列号不能为空")
    if value.isdigit():
        index = int(value)
        if index < 1:
            raise ValueError("列号必须大于 0")
        return index
    for cell in worksheet[header_row]:
        if str(cell.value or "").strip() == value:
            return cell.column
    if re.fullmatch(r"[A-Za-z]{1,3}", value):
        return column_index_from_string(value.upper())
    raise ValueError(f"第 {header_row} 行中找不到列标题：{value}")


def json_rows_fill_template_xlsx(
    inputs: dict[str, Path], output_dir: Path, params: dict[str, Any]
) -> ProcessorResult:
    records = _require_json_array(_input_path(inputs, "source_json"))
    if any(not isinstance(record, dict) for record in records):
        raise TypeError("填充模板的 JSON 数组元素必须全部是对象")
    workbook = load_workbook(_input_path(inputs, "template_xlsx"))
    sheet_name = str(params.get("sheet_name", "")).strip()
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"模板中不存在工作表：{sheet_name}")
        worksheet = workbook[sheet_name]
    else:
        worksheet = workbook.active
    header_row = int(params.get("header_row", 1))
    start_row = int(params.get("start_row", 2))
    target_column = _resolve_column(worksheet, str(params.get("target_column", "B")), header_row)
    id_selector = str(params.get("id_column", "A")).strip()
    id_column = _resolve_column(worksheet, id_selector, header_row) if id_selector else None
    id_field = str(params.get("id_field", "id"))
    value_field = str(params.get("value_field", "json_body"))
    preview: list[dict[str, Any]] = []
    for offset, record in enumerate(records):
        row_number = start_row + offset
        if id_column is not None:
            worksheet.cell(
                row=row_number,
                column=id_column,
                value=_excel_safe(record.get(id_field)),
            )
        worksheet.cell(
            row=row_number,
            column=target_column,
            value=_excel_safe(record.get(value_field)),
        )
        if len(preview) < 50:
            preview.append({id_field: record.get(id_field), value_field: record.get(value_field)})
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / "filled_template.xlsx"
    workbook.save(output_path)
    return ProcessorResult.from_path(
        output_path,
        summary=f"已向工作表 {worksheet.title!r} 写入 {len(records)} 行",
        preview=preview,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


Processor = Callable[[dict[str, Path], Path, dict[str, Any]], ProcessorResult]

PROCESSORS: dict[str, Processor] = {
    "json_pretty": json_pretty,
    "json_compact": json_compact,
    "json_show_unicode": json_show_unicode,
    "json_escape_unicode": json_escape_unicode,
    "json_unwrap_one_layer": json_unwrap_one_layer,
    "text_remove_backslashes": text_remove_backslashes,
    "json_remove_newlines": json_remove_newlines,
    "json_remove_carriage_returns": json_remove_carriage_returns,
    "json_remove_tabs": json_remove_tabs,
    "json_remove_spaces": json_remove_spaces,
    "json_select_values": json_select_values,
    "text_lines_to_json_array": text_lines_to_json_array,
    "json_trim_string_values": json_trim_string_values,
    "json_filter_empty_values": json_filter_empty_values,
    "json_deduplicate_values": json_deduplicate_values,
    "json_values_to_lines": json_values_to_lines,
    "json_values_to_quoted_list": json_values_to_quoted_list,
    "json_render_template_rows": json_render_template_rows,
    "json_map_records": json_map_records,
    "json_rows_to_xlsx": json_rows_to_xlsx,
    "json_rows_fill_template_xlsx": json_rows_fill_template_xlsx,
}
