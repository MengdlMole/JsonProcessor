import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from workbench.processors import (
    compact_json_text,
    json_deduplicate_values,
    json_filter_empty_values,
    json_remove_carriage_returns,
    json_remove_newlines,
    json_remove_spaces,
    json_remove_tabs,
    json_render_template_rows,
    json_rows_to_xlsx,
    json_select_values,
    json_unwrap_one_layer,
    pretty_json_text,
    remove_json_string_character_text,
    select_values,
    text_lines_to_json_array,
    text_remove_backslashes,
    transform_json_unicode_text,
)


class ProcessorTests(unittest.TestCase):
    def test_layout_operations_preserve_string_tokens(self) -> None:
        raw = '{ "name": "\\u4e2d\\u6587", "message": "a b" }'
        pretty = pretty_json_text(raw)
        compact = compact_json_text(pretty)
        self.assertIn('"\\u4e2d\\u6587"', pretty)
        self.assertIn('"a b"', pretty)
        self.assertEqual(compact, '{"name":"\\u4e2d\\u6587","message":"a b"}')

    def test_unicode_operation_preserves_layout(self) -> None:
        raw = '{  "name": "\\u4e2d\\u6587"  }'
        shown = transform_json_unicode_text(raw, ensure_ascii=False)
        self.assertEqual(shown, '{  "name": "中文"  }')
        self.assertEqual(transform_json_unicode_text(shown, ensure_ascii=True), raw)

    def test_remove_json_string_character_text_is_atomic(self) -> None:
        raw = '{"message":"a\\nb\\rc\\td"}'
        self.assertEqual(
            json.loads(remove_json_string_character_text(raw, "\n")),
            {"message": "ab\rc\td"},
        )
        self.assertEqual(
            json.loads(remove_json_string_character_text(raw, "\r")),
            {"message": "a\nbc\td"},
        )
        self.assertEqual(
            json.loads(remove_json_string_character_text(raw, "\t")),
            {"message": "a\nb\rcd"},
        )

    def test_atomic_text_cleanup_processors_do_only_their_declared_operation(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            wrapped = root / "wrapped.txt"
            wrapped.write_text('"{\\"id\\":1}"')
            unwrapped = json_unwrap_one_layer({"source_text": wrapped}, root / "unwrap", {})
            self.assertEqual(Path(unwrapped.primary_artifact.path).read_text(), '{"id":1}')

            slashes = root / "slashes.txt"
            slashes.write_text(r"{\"path\":\"C:\\temp\"}")
            removed = text_remove_backslashes({"source_text": slashes}, root / "slashes-out", {})
            self.assertEqual(Path(removed.primary_artifact.path).read_text(), '{"path":"C:temp"}')

            controls = root / "controls.json"
            controls.write_text('{"message":"a x\\nb\\rc\\td"}')
            operations = [
                (json_remove_newlines, {"message": "a xb\rc\td"}),
                (json_remove_carriage_returns, {"message": "a x\nbc\td"}),
                (json_remove_tabs, {"message": "a x\nb\rcd"}),
                (json_remove_spaces, {"message": "ax\nb\rc\td"}),
            ]
            for index, (processor, expected) in enumerate(operations):
                cleaned = processor({"source_json": controls}, root / f"controls-out-{index}", {})
                self.assertEqual(json.loads(Path(cleaned.primary_artifact.path).read_text()), expected)

    def test_select_values_supports_wildcard(self) -> None:
        data = {"data": {"items": [{"id": 1}, {"id": 2}]}}
        self.assertEqual(select_values(data, "data.items[*].id"), [1, 2])

    def test_select_values_rejects_partially_valid_path(self) -> None:
        with self.assertRaisesRegex(ValueError, "无效字段路径"):
            select_values({"items": []}, "items..id")

    def test_atomic_nodes_compose_without_hidden_transformations(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "input.json"
            source.write_text('{"items":[{"id":"A"},{"id":""},{"id":"A"},{"id":"B"}]}')
            selected = json_select_values(
                {"source_json": source}, root / "select", {"field_path": "items[*].id"}
            )
            self.assertEqual(json.loads(Path(selected.primary_artifact.path).read_text()), ["A", "", "A", "B"])
            filtered = json_filter_empty_values(
                {"source_json": Path(selected.primary_artifact.path)}, root / "filter", {}
            )
            unique = json_deduplicate_values(
                {"source_json": Path(filtered.primary_artifact.path)}, root / "dedupe", {}
            )
            self.assertEqual(json.loads(Path(unique.primary_artifact.path).read_text()), ["A", "B"])
            rendered = json_render_template_rows(
                {"source_json": Path(unique.primary_artifact.path)},
                root / "render",
                {
                    "id_variable": "id",
                    "common_params": "{}",
                    "json_template": '{"id": {{ id | tojson }}}',
                    "column_name": "body",
                },
            )
            workbook_result = json_rows_to_xlsx(
                {"source_json": Path(rendered.primary_artifact.path)},
                root / "xlsx",
                {"sheet_name": "Data"},
            )
            sheet = load_workbook(workbook_result.primary_artifact.path).active
            self.assertEqual(sheet["A2"].value, "A")
            self.assertEqual(json.loads(sheet["B2"].value), {"id": "A"})

    def test_split_lines_does_not_trim_filter_or_deduplicate(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "ids.txt"
            source.write_text(" A \n\n A \n")
            result = text_lines_to_json_array({"source_text": source}, root / "out", {})
            self.assertEqual(json.loads(Path(result.primary_artifact.path).read_text()), [" A ", "", " A "])

    def test_xlsx_writer_escapes_formula_like_text(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "rows.json"
            source.write_text('[{"value":"=1+1"}]')
            result = json_rows_to_xlsx(
                {"source_json": source}, root / "out", {"sheet_name": "Data"}
            )
            sheet = load_workbook(result.primary_artifact.path, data_only=False).active
            self.assertEqual(sheet["A2"].value, "'=1+1")


if __name__ == "__main__":
    unittest.main()
