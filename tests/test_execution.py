import json
import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

from workbench.execution import ExecutionService
from workbench.models import (
    InputDefinition,
    OutputDefinition,
    ToolDefinition,
    WorkflowStepDefinition,
)
from workbench.registry import ToolRegistry
from workbench.storage import WorkbenchStorage

ROOT = Path(__file__).resolve().parents[1]


class ExecutionTests(unittest.TestCase):
    def test_saved_custom_workflow_is_loaded_and_executable(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            storage = WorkbenchStorage(Path(temp))
            storage.save_custom_tool(
                ToolDefinition(
                    id="custom_format",
                    name="Custom format",
                    kind="workflow",
                    description="Saved workflow",
                    inputs=[
                        InputDefinition(id="source_json", label="JSON", content_type="json")
                    ],
                    outputs=[OutputDefinition(content_type="json")],
                    steps=[
                        WorkflowStepDefinition(
                            id="step_1",
                            uses="json_pretty",
                            inputs={"source_json": "${inputs.source_json}"},
                            params={},
                        )
                    ],
                )
            )
            registry = ToolRegistry([ROOT / "registry" / "tools.yaml", storage.custom_tools_path])
            execution = ExecutionService(registry, storage)
            _, result = execution.run_tool(
                registry.get("custom_format"),
                {"source_json": ("input.json", b'{"id":1}')},
                {},
            )
            self.assertEqual(Path(result.primary_artifact.path).read_text(), '{\n  "id": 1\n}')

    def test_atomic_workflow_fills_existing_template(self) -> None:
        registry = ToolRegistry(ROOT / "registry" / "tools.yaml")
        with tempfile.TemporaryDirectory() as temp:
            storage = WorkbenchStorage(Path(temp))
            execution = ExecutionService(registry, storage)
            workbook = Workbook()
            workbook.active.append(["ID", "Body", "Keep"])
            workbook.active["C2"] = "unchanged"
            template = BytesIO()
            workbook.save(template)
            _, result = execution.run_tool(
                registry.get("txt_ids_fill_template_workflow"),
                {
                    "source_text": ("ids.txt", b" A-1 \n\nA-1\nB-2\n"),
                    "template_xlsx": ("template.xlsx", template.getvalue()),
                },
                {
                    "id_variable": "id",
                    "common_params": "{}",
                    "json_template": '{"id": {{ id | tojson }}}',
                    "column_name": "json_body",
                    "sheet_name": "",
                    "header_row": 1,
                    "start_row": 2,
                    "id_column": "ID",
                    "target_column": "Body",
                },
            )
            output = load_workbook(result.primary_artifact.path).active
            self.assertEqual(output["A2"].value, "A-1")
            self.assertEqual(output["A3"].value, "B-2")
            self.assertEqual(json.loads(output["B2"].value), {"id": "A-1"})
            self.assertEqual(output["C2"].value, "unchanged")

    def test_unknown_parameter_is_rejected_before_a_run_is_created(self) -> None:
        registry = ToolRegistry(ROOT / "registry" / "tools.yaml")
        with tempfile.TemporaryDirectory() as temp:
            storage = WorkbenchStorage(Path(temp))
            execution = ExecutionService(registry, storage)
            with self.assertRaisesRegex(ValueError, "未知参数"):
                execution.run_tool(
                    registry.get("json_pretty"),
                    {"source_json": ("input.json", b'{"id": 1}')},
                    {"unexpected": True},
                )
            self.assertEqual(storage.list_runs(), [])

    def test_declarative_workflow_records_steps_and_passes_artifact(self) -> None:
        registry = ToolRegistry(ROOT / "registry" / "tools.yaml")
        with tempfile.TemporaryDirectory() as temp:
            storage = WorkbenchStorage(Path(temp))
            execution = ExecutionService(registry, storage)
            tool = registry.get("json_ids_to_json_xlsx")
            input_bytes = json.dumps(
                {"data": {"items": [{"id": "A-1"}, {"id": "B-2"}]}},
                ensure_ascii=False,
            ).encode()
            run_id, result = execution.run_tool(
                tool,
                {"source_json": ("input.json", input_bytes)},
                {
                    "field_path": "data.items[*].id",
                    "id_variable": "id",
                    "common_params": '{"source":"workflow"}',
                    "json_template": ('{"id": {{ id | tojson }}, "source": {{ source | tojson }}}'),
                    "column_name": "json_body",
                    "sheet_name": "Data",
                },
            )

            self.assertIn("5/5", result.summary)
            self.assertEqual(
                [step["step_id"] for step in storage.list_step_runs(run_id)],
                ["extract", "filter_empty", "deduplicate", "render", "write_xlsx"],
            )
            workbook = load_workbook(result.primary_artifact.path)
            self.assertEqual(workbook.active["A2"].value, "A-1")
            self.assertEqual(
                json.loads(workbook.active["B2"].value),
                {"id": "A-1", "source": "workflow"},
            )
            manifest = json.loads((Path(temp) / "runs" / run_id / "run.json").read_text())
            self.assertEqual(manifest["status"], "success")
            self.assertEqual(
                [step["status"] for step in manifest["steps"]],
                ["success", "success", "success", "success", "success"],
            )

    def test_failed_script_is_persisted(self) -> None:
        registry = ToolRegistry(ROOT / "registry" / "tools.yaml")
        with tempfile.TemporaryDirectory() as temp:
            storage = WorkbenchStorage(Path(temp))
            execution = ExecutionService(registry, storage)
            with self.assertRaises(ValueError):
                execution.run_tool(
                    registry.get("json_select_values"),
                    {"source_json": ("input.json", b'{"items": []}')},
                    {"field_path": "items..id"},
                )
            run = storage.list_runs(1)[0]
            self.assertEqual(run["status"], "failed")
            self.assertIn("无效字段路径", run["error"])
            steps = storage.list_step_runs(run["id"])
            self.assertEqual(steps[0]["status"], "failed")


if __name__ == "__main__":
    unittest.main()
